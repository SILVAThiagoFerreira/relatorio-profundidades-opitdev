"""Workbook reading layer."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from .exceptions import DataReadError
from .models import AppConfig, RawDataset, RawRow, WorkbookSpec


def _normalize_header(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _resolve_input_path(config: AppConfig, spec: WorkbookSpec) -> Path:
    path = Path(spec.file_name)
    return path if path.is_absolute() else (config.input_dir / path).resolve()


def read_workbook(config: AppConfig, spec: WorkbookSpec) -> RawDataset:
    """Read a configured workbook into a raw row snapshot."""

    path = _resolve_input_path(config, spec)
    if not path.exists():
        raise DataReadError(f"Input file not found: {path}")

    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:  # pragma: no cover - openpyxl raises many concrete errors
        raise DataReadError(f"Unable to read workbook: {path}") from exc

    try:
        sheet_name = spec.sheet_name or workbook.sheetnames[0]
        if sheet_name not in workbook.sheetnames:
            raise DataReadError(f"Sheet '{sheet_name}' not found in workbook '{path.name}'")

        worksheet = workbook[sheet_name]
        try:
            header_values = next(worksheet.iter_rows(min_row=spec.header_row, max_row=spec.header_row, values_only=True))
        except StopIteration as exc:
            raise DataReadError(f"Header row {spec.header_row} not found in sheet '{sheet_name}'") from exc

        headers = [_normalize_header(value) for value in header_values]
        if not any(headers):
            raise DataReadError(f"Header row {spec.header_row} is empty in sheet '{sheet_name}'")

        normalized_headers: list[str] = []
        seen_headers: set[str] = set()
        for header in headers:
            if not header:
                normalized_headers.append("")
                continue
            if header in seen_headers:
                raise DataReadError(f"Duplicate header '{header}' found in sheet '{sheet_name}'")
            seen_headers.add(header)
            normalized_headers.append(header)

        rows: list[RawRow] = []
        for row_number, values in enumerate(worksheet.iter_rows(min_row=spec.header_row + 1, values_only=True), start=spec.header_row + 1):
            row_values: dict[str, object] = {}
            has_data = False
            for index, header in enumerate(normalized_headers):
                if not header:
                    continue
                value = values[index] if index < len(values) else None
                if value not in (None, ""):
                    has_data = True
                row_values[header] = value
            if has_data:
                rows.append(RawRow(row_number=row_number, values=row_values))

        return RawDataset(label=spec.label, path=path, sheet_name=sheet_name, headers=[header for header in normalized_headers if header], rows=rows)
    finally:
        workbook.close()
